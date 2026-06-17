from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    extension: str
    text: str
    word_count: int
    char_count: int


def extract_document(filename: str, content: bytes) -> ExtractedDocument:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type '{extension}'. Supported: {supported}.")

    if extension == ".pdf":
        text = extract_pdf(content)
    elif extension == ".docx":
        text = extract_docx(content)
    else:
        text = decode_text(content)

    cleaned = normalize_text(text)
    return ExtractedDocument(
        filename=filename,
        extension=extension,
        text=cleaned,
        word_count=len(cleaned.split()),
        char_count=len(cleaned),
    )


def extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    table_lines: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_lines.append(" | ".join(cells))
    return "\n\n".join(paragraphs + table_lines)


def decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def normalize_text(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    collapsed: list[str] = []
    blank = False
    for line in lines:
        if not line.strip():
            if not blank:
                collapsed.append("")
            blank = True
            continue
        collapsed.append(line.strip())
        blank = False
    return "\n".join(collapsed).strip()
