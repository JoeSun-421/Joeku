from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_PARAGRAPH_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from academic_agent.export_options import CoverPageOptions, ExportFormatOptions


def markdown_to_docx(
    markdown_text: str,
    output_path: Path,
    *,
    title: str | None = None,
    options: ExportFormatOptions | None = None,
) -> Path:
    """Convert academic paper markdown to a Word (.docx) document."""
    opts = options or ExportFormatOptions()
    family = opts.resolved_family()["docx"]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = family
    style.font.size = Pt(opts.font_size)
    if style._element.rPr is not None:
        style._element.rPr.rFonts.set(qn("w:eastAsia"), family)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    style.paragraph_format.line_spacing = opts.line_spacing

    if opts.cover_page and any(getattr(opts.cover_page, f, "") for f in CoverPageOptions.model_fields):
        _add_cover_page(doc, opts.cover_page, opts, family)

    lines = markdown_text.splitlines()
    idx = 0
    doc_title = title
    in_references = False

    while idx < len(lines):
        line = lines[idx].rstrip()
        idx += 1

        if not line.strip():
            continue

        if line.startswith("# ") and doc_title is None:
            doc_title = line[2:].strip()
            p = doc.add_heading(doc_title, level=0)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            _apply_para_format(p, opts, family, heading_scale=1.35)
            continue

        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
            _apply_para_format(p, opts, family, heading_scale=1.25)
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            in_references = heading.lower() in {
                "references",
                "works cited",
                "bibliography",
                "verification checklist",
                "参考文献",
            }
            p = doc.add_heading(heading, level=2)
            _apply_para_format(p, opts, family, heading_scale=1.12)
            continue

        if line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=3)
            _apply_para_format(p, opts, family, heading_scale=1.05)
            continue

        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_formatted(p, line[2:].strip(), family=family, font_size=opts.font_size)
            _apply_para_format(p, opts, family, body_indent=False)
            continue

        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            _add_inline_formatted(p, re.sub(r"^\d+\.\s", "", line).strip(), family=family, font_size=opts.font_size)
            _apply_para_format(p, opts, family, body_indent=False)
            continue

        text = line.strip()
        if text.startswith("**") and text.endswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(text.strip("*"))
            run.bold = True
            _style_run(run, family, opts.font_size)
            _apply_para_format(p, opts, family, body_indent=False)
            continue

        p = doc.add_paragraph()
        _add_inline_formatted(p, line, family=family, font_size=opts.font_size)
        _apply_para_format(p, opts, family, body_indent=not in_references)

    margin = Inches(opts.margin_inches)
    for section in doc.sections:
        section.top_margin = margin
        section.bottom_margin = margin
        section.left_margin = margin
        section.right_margin = margin

    output_path = Path(output_path)
    if output_path.suffix.lower() not in {".docx", ".doc"}:
        output_path = output_path.with_suffix(".docx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _add_cover_page(doc: Document, cover: CoverPageOptions, opts: ExportFormatOptions, family: str) -> None:
    doc.add_paragraph()
    for _ in range(4):
        doc.add_paragraph()

    def _center_line(text: str, *, size_scale: float = 1.0, bold: bool = False) -> None:
        if not text.strip():
            return
        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = p.add_run(text.strip())
        run.bold = bold
        _style_run(run, family, int(opts.font_size * size_scale))
        _apply_para_format(p, opts, family, heading_scale=size_scale, body_indent=False)

    _center_line(cover.title or "论文标题", size_scale=1.8, bold=True)
    _center_line(cover.subtitle, size_scale=1.15)
    doc.add_paragraph()
    _center_line(cover.author)
    _center_line(cover.institution)
    _center_line(cover.course)
    _center_line(cover.instructor)
    _center_line(cover.date)
    doc.add_page_break()


def _apply_para_format(
    paragraph,
    opts: ExportFormatOptions,
    family: str,
    *,
    heading_scale: float = 1.0,
    body_indent: bool = False,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = opts.line_spacing
    if body_indent and opts.first_line_indent_chars > 0:
        fmt.first_line_indent = Pt(opts.font_size * opts.first_line_indent_chars)
    for run in paragraph.runs:
        _style_run(run, family, int(opts.font_size * heading_scale))


def _style_run(run, family: str, size: int) -> None:
    run.font.name = family
    run.font.size = Pt(size)
    if run._element.rPr is not None:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), family)


def _add_hyperlink(paragraph, text: str, url: str, *, family: str, font_size: int) -> None:
    if not url:
        run = paragraph.add_run(text)
        _style_run(run, family, font_size)
        return
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), family)
    r_fonts.set(qn("w:hAnsi"), family)
    r_fonts.set(qn("w:eastAsia"), family)
    r_pr.append(r_fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(font_size * 2)))
    r_pr.append(sz)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


_ET_AL_RE = re.compile(r"\bet\s+al\.?", re.IGNORECASE)
_INLINE_MARKDOWN_RE = re.compile(
    r"(\*\*[^*\n]+?\*\*|\*[^*\n]+?\*|`[^`\n]+`|\[[^\]]+\]\(https?://[^)]+\))"
)


def _normalize_et_al(text: str) -> str:
    return _ET_AL_RE.sub("et\u00a0al.", text)


def _unwrap_markdown_spans(text: str) -> str:
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    return text


def _unwrap_markdown_in_parenthetical_citations(text: str) -> str:
    def _clean(match: re.Match[str]) -> str:
        return f"({_unwrap_markdown_spans(match.group(1))})"

    return re.sub(r"\(([^)]*\d{4}[a-z]?[^)]*)\)", _clean, text)


def _unwrap_et_al_markdown(text: str) -> str:
    text = re.sub(r"\*\*([^*]*et\s+al\.?[^*]*)\*\*", r"\1", text, flags=re.IGNORECASE)
    text = re.sub(r"(?<!\*)\*([^*]*et\s+al\.?[^*]*)\*(?!\*)", r"\1", text, flags=re.IGNORECASE)
    return text


def _prepare_inline_text(text: str) -> str:
    text = _unwrap_et_al_markdown(text)
    text = _unwrap_markdown_in_parenthetical_citations(text)
    return _normalize_et_al(text)


def _strip_markdown_inline(text: str) -> str:
    text = re.sub(r"\[(.+?)\]\((https?://[^)]+)\)", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text


def _add_inline_formatted(paragraph, text: str, *, family: str, font_size: int) -> None:
    text = _prepare_inline_text(text)
    pos = 0
    for match in _INLINE_MARKDOWN_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(_strip_markdown_inline(text[pos : match.start()]))
            _style_run(run, family, font_size)
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            _style_run(run, family, font_size)
        elif token.startswith("*") and token.endswith("*") and not token.startswith("**"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            _style_run(run, family, font_size)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(font_size)
        elif token.startswith("["):
            m = re.match(r"\[(.+?)\]\((https?://[^)]+)\)", token)
            if m:
                _add_hyperlink(paragraph, m.group(1), m.group(2), family=family, font_size=font_size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(_strip_markdown_inline(text[pos:]))
        _style_run(run, family, font_size)